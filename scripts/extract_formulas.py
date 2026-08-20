"""Dump header row + first data-row formulas from MODEL.xlsx sheets.

The output in docs/formulas/ is the spec for the Python port: one JSON file
per sheet mapping column header -> formula (or literal) from the first data row.
"""

import json
import sys
from pathlib import Path

import openpyxl

WORKBOOK = Path.home() / "Downloads" / "MODEL.xlsx"
OUT = Path(__file__).resolve().parents[1] / "docs" / "formulas"

# sheet -> (header_row, data_row, max_col)
SHEETS = {
    "CONTROL": None,  # full dump, small sheet
    "PREDICT": (2, 3, 60),
    "MODEL": (1, 2, 250),
    "PPts": (1, 2, 70),
    "TABLES": (1, 2, 80),
    "TABLE": (1, 2, 30),
    "PSTABLE": (1, 2, 30),
    "FIXTURES": (1, 2, 80),
    "TEAMS": (1, 2, 25),
}


def cell_repr(c):
    v = c.value
    if v is None:
        return None
    if hasattr(v, "text"):  # ArrayFormula
        return f"{{ARRAY}}{v.text}"
    return v


def main():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=False)
    OUT.mkdir(parents=True, exist_ok=True)

    for name, spec in SHEETS.items():
        ws = wb[name]
        if spec is None:
            rows = []
            for row in ws.iter_rows(min_row=1, max_row=130, max_col=10):
                for c in row:
                    v = cell_repr(c)
                    if v is not None:
                        rows.append({"cell": c.coordinate, "value": str(v)})
            (OUT / f"{name}.json").write_text(json.dumps(rows, indent=1))
            print(f"{name}: {len(rows)} cells")
            continue

        hrow, drow, maxcol = spec
        headers, formulas = {}, {}
        for row in ws.iter_rows(min_row=hrow, max_row=hrow, max_col=maxcol):
            for c in row:
                if c.value is not None:
                    headers[c.column_letter] = str(c.value)
        for row in ws.iter_rows(min_row=drow, max_row=drow, max_col=maxcol):
            for c in row:
                v = cell_repr(c)
                if v is not None:
                    formulas[c.column_letter] = str(v)
        out = {
            "sheet": name,
            "header_row": hrow,
            "data_row": drow,
            "columns": {
                col: {"header": headers.get(col), "formula": formulas.get(col)}
                for col in sorted(set(headers) | set(formulas), key=lambda s: (len(s), s))
            },
        }
        (OUT / f"{name}.json").write_text(json.dumps(out, indent=1))
        print(f"{name}: {len(out['columns'])} columns")

    wb.close()


if __name__ == "__main__":
    sys.exit(main())
