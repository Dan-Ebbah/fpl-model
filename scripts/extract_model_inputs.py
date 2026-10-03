"""Extract the MODEL sheet's remaining input data from MODEL.xlsx.

The MODEL-sheet builder (fpl_model.model_build) needs four inputs that earlier
extraction passes didn't capture:

- PT / PSPT       — FBRef "Playing Time" tables pasted into the workbook
                    (current and previous season): Starts/Subs/unSub/minutes.
- MINS/xG/xA/DC   — per-gameweek delta matrices (rows 1001-1999 of the pivot
                    sheets): player code -> value per GW-snapshot column.
- CONTROL values  — the computed cell values of the CONTROL sheet (the earlier
                    extract captured formulas; the builder needs the numbers).

All are frozen last-season data (or static config) and change at most once a
season, so they live in data/reference/ alongside the other extracts.

Usage: uv run python scripts/extract_model_inputs.py [path-to-MODEL.xlsx]
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd

WORKBOOK = Path(sys.argv[1] if len(sys.argv) > 1 else "~/Downloads/MODEL.xlsx").expanduser()
OUT = Path(__file__).resolve().parents[1] / "data" / "reference"

# The named tables PT (A2:Y722) and PSPT (A2:Y762): header on the first row of
# the range. Repeated FBRef page headers appear as data rows; keep them —
# Excel's MATCH scans them too and the builder reproduces that faithfully.
FBREF_TABLES = {"PT": (2, 722), "PSPT": (2, 762)}

# Pivot delta blocks: rows 1001-1999, key in col A, 103 value columns B..CZ.
PIVOT_SHEETS = ["MINS", "xG", "xA", "DC"]
PIVOT_ROWS = (1001, 1999)
PIVOT_MAX_COL = 104  # CZ


def extract_fbref(wb: openpyxl.Workbook, name: str, first_row: int, last_row: int) -> None:
    ws = wb[name]
    rows = list(ws.iter_rows(min_row=first_row, max_row=last_row, max_col=25, values_only=True))
    df = pd.DataFrame(rows[1:], columns=[str(c) for c in rows[0]])
    # mixed numeric/text columns (repeated FBRef page headers): store as text,
    # the builder coerces back to numbers
    df = df.map(lambda v: None if v is None else str(v))
    df.to_parquet(OUT / f"{name}.parquet")
    print(f"{name}: {len(df)} rows -> {name}.parquet")


def extract_pivot(wb: openpyxl.Workbook, name: str) -> None:
    ws = wb[name]
    rows = list(
        ws.iter_rows(
            min_row=PIVOT_ROWS[0], max_row=PIVOT_ROWS[1], max_col=PIVOT_MAX_COL, values_only=True
        )
    )
    # Positional column names (header labels can repeat); the GW labels sit in
    # the block's first row (sheet row 1001), which stays as data row 0 so the
    # builder can reproduce Excel's MATCH over the whole 1001-1999 range.
    cols = ["key"] + [f"c{i}" for i in range(1, PIVOT_MAX_COL)]
    df = pd.DataFrame(rows, columns=cols)
    for c in cols:
        df[c] = df[c].map(lambda v: None if v is None else str(v))
    df.to_parquet(OUT / f"{name}_delta.parquet")
    labels = [v for v in df.iloc[0, 1:] if v is not None]
    print(f"{name}: {len(df)} rows, labels {labels[0]}..{labels[-1]} -> {name}_delta.parquet")


def extract_control(wb: openpyxl.Workbook) -> None:
    ws = wb["CONTROL"]
    records = []
    for row in ws.iter_rows(min_row=1, max_row=210, max_col=8):
        for cell in row:
            if cell.value is not None:
                records.append({"cell": cell.coordinate, "value": str(cell.value)})
    df = pd.DataFrame(records)
    df.to_parquet(OUT / "CONTROL_values.parquet")
    print(f"CONTROL: {len(df)} cells -> CONTROL_values.parquet")


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    extract_control(wb)
    for name, (first, last) in FBREF_TABLES.items():
        extract_fbref(wb, name, first, last)
    for name in PIVOT_SHEETS:
        extract_pivot(wb, name)
    wb.close()


if __name__ == "__main__":
    main()
